"""Keep Verdict/Verdicts.xlsx in sync with the <grievance_id>Verdict.json files.

New grievances are appended as rows; a grievance that is already in the sheet
(e.g. re-checked with --force) has its row updated in place. Rows already in
the sheet are never removed, so the workbook keeps growing across runs.

Used by app.compare_atr: each verdict is written to the sheet as soon as it
is produced, and the whole Verdict folder is re-synced at the end of the run.
"""

import json
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


WORKBOOK_NAME = "Verdicts.xlsx"
SHEET = "Verdicts"
SUMMARY_SHEET = "Summary"

# (header, verdict key, column width)
COLUMNS = [
    ("Case ID", "case_id", 12),
    ("Verdict", "verdict", 22),
    ("Resolved?", "is_resolved", 11),
    ("Complaint Summary", "complaint_summary", 50),
    ("Final Action", "final_action", 50),
    ("Reason", "reason", 60),
    ("Formal Remark", "formal_remark", 45),
    ("Verdict Meaning", "verdict_meaning", 45),
    ("Model", "model", 20),
    ("Prompt Tokens", "prompt_tokens", 14),
    ("Output Tokens", "output_tokens", 14),
]

VERDICT_COLOURS = {
    "RESOLVED": "C6EFCE",
    "PARTIALLY_RESOLVED": "FFEB9C",
    "REDIRECTED": "DDEBF7",
    "REJECTED": "F8CBAD",
    "CLOSED_WITHOUT_ACTION": "D9D2E9",
    "WRONG_ACTION": "FFC7CE",
    "UNCLEAR": "E7E6E6",
}

_thin = Side(style="thin", color="BFBFBF")
BORDER = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(bold=True, color="FFFFFF")


def _style_header(cell):
    cell.fill = HEADER_FILL
    cell.font = HEADER_FONT
    cell.border = BORDER
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _new_verdict_sheet(wb):
    ws = wb.create_sheet(SHEET, 0)
    for col, (header, _, width) in enumerate(COLUMNS, start=1):
        _style_header(ws.cell(1, col, header))
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.freeze_panes = "B2"
    return ws


def _cell_value(key, verdict):
    value = verdict.get(key)
    if key == "is_resolved":
        return "Yes" if value else "No"
    if key == "case_id" and str(value).isdigit():
        return int(value)
    return value


def _write_row(ws, row, verdict):
    for col, (_, key, _) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row, col, _cell_value(key, verdict))
        cell.border = BORDER
        cell.alignment = Alignment(vertical="top", wrap_text=True)
    verdict_cell = ws.cell(row, 2)
    verdict_cell.fill = PatternFill("solid", fgColor=VERDICT_COLOURS.get(verdict.get("verdict"), "FFFFFF"))
    verdict_cell.font = Font(bold=True)


def _rebuild_summary(wb):
    if SUMMARY_SHEET in wb.sheetnames:
        del wb[SUMMARY_SHEET]
    s = wb.create_sheet(SUMMARY_SHEET)
    for col, header in enumerate(("Verdict", "Count", "Share"), start=1):
        _style_header(s.cell(1, col, header))

    total_row = len(VERDICT_COLOURS) + 2
    for row, (verdict, colour) in enumerate(VERDICT_COLOURS.items(), start=2):
        s.cell(row, 1, verdict).fill = PatternFill("solid", fgColor=colour)
        s.cell(row, 2, f"=COUNTIF({SHEET}!$B:$B,A{row})")
        s.cell(row, 3, f"=IF($B${total_row}=0,0,B{row}/$B${total_row})").number_format = "0%"
    s.cell(total_row, 1, "Total").font = Font(bold=True)
    s.cell(total_row, 2, f"=SUM(B2:B{total_row - 1})").font = Font(bold=True)
    s.cell(total_row, 3, f"=IF($B${total_row}=0,0,1)").number_format = "0%"

    for row in range(2, total_row + 1):
        for col in (1, 2, 3):
            s.cell(row, col).border = BORDER
    s.column_dimensions["A"].width = 26
    s.column_dimensions["B"].width = 10
    s.column_dimensions["C"].width = 10


def update_workbook(verdicts, workbook_path: Path) -> tuple[int, int]:
    """Append new verdicts / update existing ones. Returns (added, updated)."""
    if workbook_path.exists():
        wb = load_workbook(workbook_path)
        ws = wb[SHEET] if SHEET in wb.sheetnames else _new_verdict_sheet(wb)
    else:
        wb = Workbook()
        wb.remove(wb.active)
        ws = _new_verdict_sheet(wb)

    existing_rows = {
        str(ws.cell(row, 1).value): row
        for row in range(2, ws.max_row + 1)
        if ws.cell(row, 1).value is not None
    }

    added = updated = 0
    for verdict in verdicts:
        case_id = str(verdict.get("case_id"))
        row = existing_rows.get(case_id)
        if row is None:
            row = ws.max_row + 1
            existing_rows[case_id] = row
            added += 1
        else:
            updated += 1
        _write_row(ws, row, verdict)

    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{max(ws.max_row, 2)}"
    _rebuild_summary(wb)

    try:
        wb.save(workbook_path)
    except PermissionError:
        raise PermissionError(
            f"Could not write {workbook_path} - close it in Excel and run again."
        ) from None
    return added, updated


def sync_verdict_dir(verdict_dir: Path) -> Path:
    verdicts = [
        json.loads(path.read_text(encoding="utf-8"))
        for path in sorted(verdict_dir.glob("*Verdict.json"))
    ]
    workbook_path = verdict_dir / WORKBOOK_NAME
    added, updated = update_workbook(verdicts, workbook_path)
    print(f"Excel: {added} new row(s), {updated} refreshed -> {workbook_path}")
    return workbook_path

