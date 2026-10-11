"""Check every grievance in pgrs30_txt with Gemma or Qwen and compare the verdicts with
the reference batch in Observation/Observation1/Observation1.xlsx.

Reads:
    pgrs30_txt/<id>/ATR/*.txt, Citizen/*.txt       (or, with --old-ocr, PGRS30_Out/<id>/ATR.txt, Citizen.txt)
Writes only one Excel file per run (no JSON files):
    Observation/<folder>/<folder>.xlsx   sheets Observations, Summary and Comparison (this run next to Observation1)
Each observation is added to the workbook as soon as it is ready, so an interrupted
run keeps what it finished; a grievance already in the workbook is skipped next time.

Usage:
    python compare_pgrs30.py                   # Gemma 4 E4B on the remote server
    python compare_pgrs30.py --model gemma     # Gemma 4 E2B on local Ollama
    python compare_pgrs30.py --model qwen14b   # Qwen3 14B on the remote server
    python compare_pgrs30.py --model qwen14b --old-ocr   # same, on the earlier OCR in PGRS30_Out
    python compare_pgrs30.py --force           # re-check grievances already in the workbook
"""

import argparse
import sys
import time
from concurrent.futures import ThreadPoolExecutor, as_completed
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill

from app.excel import COLUMNS, RATING_STARS, SHEET, _style_header, update_workbook
from app.observation import MODELS, OBSERVATION_ROOTS, PROJECT_ROOT, check_case


TEXT_DIR = PROJECT_ROOT / "pgrs30_txt"
OLD_TEXT_DIR = PROJECT_ROOT / "PGRS30_Out"  # the earlier OCR: <id>/ATR.txt, Citizen.txt
REFERENCE = next(
    (p for p in (OBSERVATION_ROOTS[0] / "Observation1" / "Observation1.xlsx",
                 PROJECT_ROOT / "Observation1" / "Observation1.xlsx") if p.exists()),
    OBSERVATION_ROOTS[0] / "Observation1" / "Observation1.xlsx",
)
COMPARISON_SHEET = "Comparison"
# model -> (output folder, short name, grievances checked at the same time)
RUNS = {
    "gemma-e4b": ("PGRS30_Gemma4_E4B", "E4B", 3),  # remote server
    "gemma": ("PGRS30_Gemma4_E2B", "E2B", 1),      # local Ollama runs one at a time
    "qwen14b": ("PGRS30_Qwen3_14B", "Qwen14B", 2),  # remote server
}


def read_texts(case_dir: Path, subfolder: str) -> str:
    """pgrs30_txt keeps one .txt per document in ATR/ and Citizen/; PGRS30_Out one ATR.txt, Citizen.txt."""
    files = sorted((case_dir / subfolder).glob("*.txt")) or [
        p for p in [case_dir / f"{subfolder}.txt"] if p.exists()
    ]
    return "\n\n".join(p.read_text(encoding="utf-8") for p in files).strip()


def read_workbook(path: Path) -> dict[str, dict]:
    """Case ID -> observation, read back from an observation workbook ("Yes" -> True, "★★" -> 2)."""
    if not path.exists():
        return {}
    observations = {}
    rows = load_workbook(path, read_only=True)[SHEET].iter_rows(values_only=True)
    # Match columns by header, since Observation1.xlsx has no Rating column.
    keys = {header: key for header, key, _ in COLUMNS}
    header = [keys.get(h) for h in next(rows)]
    for row in rows:
        if row[0] is None or row[1] not in ("Yes", "No"):  # e.g. a "Not checked" placeholder row
            continue
        observation = {key: value for key, value in zip(header, row) if key}
        observation = {key: observation.get(key) for _, key, _ in COLUMNS}
        observation["case_id"] = str(observation["case_id"])
        observation["is_resolved"] = observation["is_resolved"] == "Yes"
        observation["rating"] = str(observation["rating"] or "").count(RATING_STARS)
        observations[observation["case_id"]] = observation
    return observations


def check(case_dir: Path, model: str) -> tuple[str, dict | None, str]:
    citizen, atr = read_texts(case_dir, "Citizen"), read_texts(case_dir, "ATR")
    missing = [name for name, text in (("Citizen", citizen), ("ATR", atr)) if not text]
    if missing:
        return case_dir.name, None, f"skipped (no {' or '.join(missing)} text)"
    started = time.time()
    try:
        observation = check_case(case_dir.name, citizen, atr, model)
    except Exception as exc:
        return case_dir.name, None, f"FAILED: {exc}"
    status = "Resolved" if observation["is_resolved"] else "Not resolved"
    return case_dir.name, observation, f"{status}, rating {observation['rating']} ({time.time() - started:.0f}s)"


def write_comparison(workbook: Path, name: str, case_ids: list[str], notes: dict[str, str]):
    """Add (or rebuild) the Comparison sheet: this run next to Observation1.xlsx."""
    new, ref = read_workbook(workbook), read_workbook(REFERENCE)
    wb = load_workbook(workbook)
    if COMPARISON_SHEET in wb.sheetnames:
        del wb[COMPARISON_SHEET]
    ws = wb.create_sheet(COMPARISON_SHEET)
    ws.append([
        "Grievance", f"{name} resolved", "Observation1 resolved", f"{name} rating", "Observation1 rating",
        "Agree", f"{name} scheme", "Observation1 scheme", f"{name} final action", "Observation1 final action",
        "Note",
    ])
    for cell in ws[1]:
        _style_header(cell)
    differ = PatternFill("solid", fgColor="FCE4D6")
    agree = total = 0
    for case_id in case_ids:
        a, b = new.get(case_id, {}), ref.get(case_id, {})
        same = None
        if a and b:
            same = a["is_resolved"] == b["is_resolved"]
            total += 1
            agree += same
        yes_no = {True: "Yes", False: "No", None: ""}
        ws.append([
            int(case_id), yes_no[a.get("is_resolved")], yes_no[b.get("is_resolved")],
            RATING_STARS * a.get("rating", 0), RATING_STARS * b.get("rating", 0),
            {True: "yes", False: "NO", None: ""}[same],
            a.get("scheme_english"), b.get("scheme_english"),
            a.get("final_action_taken"), b.get("final_action_taken"), notes.get(case_id, ""),
        ])
        if same is False:
            for cell in ws[ws.max_row]:
                cell.fill = differ
    ws.append([])
    ws.append([f"Agrees with Observation1 on {agree} of {total} grievance(s) checked by both"])
    ws.cell(ws.max_row, 1).font = Font(bold=True)
    for column, width in zip("ABCDEFGHIJK", (11, 12, 12, 10, 10, 7, 30, 30, 60, 60, 30)):
        ws.column_dimensions[column].width = width
    ws.freeze_panes = "B2"
    wb.save(workbook)
    return agree, total


def main():
    parser = argparse.ArgumentParser(description="Check pgrs30_txt with Gemma or Qwen and compare with Observation1.")
    parser.add_argument("--model", choices=RUNS, default="gemma-e4b",
                        help="gemma-e4b (default), gemma (local E2B) or qwen14b.")
    parser.add_argument("--old-ocr", action="store_true",
                        help=f"Read the earlier OCR in {OLD_TEXT_DIR.name} instead of {TEXT_DIR.name}.")
    parser.add_argument("--suffix", default="", help="Added to the output folder name, e.g. _v2, to keep an earlier run.")
    parser.add_argument("--force", action="store_true", help="Re-check grievances already in the workbook.")
    args = parser.parse_args()
    for stream in (sys.stdout, sys.stderr):
        stream.reconfigure(encoding="utf-8", errors="replace")

    folder, name, workers = RUNS[args.model]
    text_dir = OLD_TEXT_DIR if args.old_ocr else TEXT_DIR
    if args.old_ocr:
        folder += "_PGRS30_Out"
    folder += args.suffix
    out_dir = OBSERVATION_ROOTS[0] / folder
    out_dir.mkdir(parents=True, exist_ok=True)
    workbook = out_dir / f"{folder}.xlsx"

    case_dirs = sorted(p for p in text_dir.iterdir() if p.is_dir())
    done = set() if args.force else set(read_workbook(workbook))
    todo = [d for d in case_dirs if d.name not in done]
    print(f"{len(case_dirs)} grievance(s) in {text_dir.name}, {len(done)} already in {workbook.name}, "
          f"checking {len(todo)} with {MODELS[args.model]}")

    notes = {}
    with ThreadPoolExecutor(workers) as pool:
        jobs = [pool.submit(check, d, args.model) for d in todo]
        for index, job in enumerate(as_completed(jobs), start=1):
            case_id, observation, note = job.result()
            notes[case_id] = note
            if observation:
                update_workbook([observation], workbook)  # saved straight away
            print(f"[{index}/{len(jobs)}] {case_id} -> {note}", flush=True)

    if not workbook.exists():
        raise SystemExit("Nothing was checked, so no workbook was written.")
    agree, total = write_comparison(workbook, name, [d.name for d in case_dirs], notes)
    observations = read_workbook(workbook).values()
    resolved = sum(o["is_resolved"] for o in observations)
    print(f"\nResolved: {resolved}, not resolved: {len(observations) - resolved}")
    print(f"Agrees with Observation1 on {agree} of {total} grievance(s) checked by both")
    print(f"Saved to {workbook}")


if __name__ == "__main__":
    main()
